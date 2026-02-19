import io

import pytest

from src.ingestion.parser import parse_file, SUPPORTED_TYPES


def test_unsupported_file_type_raises():
    with pytest.raises(ValueError, match="Unsupported file type"):
        parse_file(b"data", "file.txt")


def test_supported_types():
    assert SUPPORTED_TYPES == {"pdf", "docx", "pptx", "xlsx"}


def test_parse_docx():
    from docx import Document

    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("First paragraph")
    doc.add_paragraph("Second paragraph")
    doc.save(buf)
    buf.seek(0)

    text = parse_file(buf.read(), "test.docx")
    assert "First paragraph" in text
    assert "Second paragraph" in text


def test_parse_xlsx():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Value"])
    ws.append(["Alice", 42])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    text = parse_file(buf.read(), "test.xlsx")
    assert "Name" in text
    assert "Alice" in text
    assert "42" in text


def test_parse_pptx():
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # title + content
    slide.shapes.title.text = "Slide Title"
    slide.placeholders[1].text = "Slide body text"
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)

    text = parse_file(buf.read(), "test.pptx")
    assert "Slide Title" in text
    assert "Slide body text" in text


def test_parse_pdf():
    from PyPDF2 import PdfWriter

    writer = PdfWriter()
    # PdfWriter alone can't easily create text; use a blank page annotation approach
    # For a minimal test, create a PDF with reportlab if available, otherwise skip
    # Simpler: just test the dispatch works with a minimal PDF
    buf = io.BytesIO()
    writer.add_blank_page(width=200, height=200)
    writer.write(buf)
    buf.seek(0)

    # Blank PDF returns empty text — parser should handle gracefully
    text = parse_file(buf.read(), "blank.pdf")
    assert isinstance(text, str)
