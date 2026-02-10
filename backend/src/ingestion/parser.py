from __future__ import annotations

import io
import logging
from pathlib import Path

logger = logging.getLogger(__name__)

SUPPORTED_TYPES = {"pdf", "docx", "pptx", "xlsx"}


def parse_file(file_bytes: bytes, filename: str) -> str:
    """Extract text from a document file. Returns raw text."""
    ext = Path(filename).suffix.lower().lstrip(".")
    if ext not in SUPPORTED_TYPES:
        raise ValueError(f"Unsupported file type: {ext}. Supported: {SUPPORTED_TYPES}")

    parsers = {
        "pdf": _parse_pdf,
        "docx": _parse_docx,
        "pptx": _parse_pptx,
        "xlsx": _parse_xlsx,
    }
    text = parsers[ext](file_bytes)
    logger.info(f"Parsed {filename}: {len(text)} chars")
    return text


def _parse_pdf(data: bytes) -> str:
    from PyPDF2 import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text)
    return "\n\n".join(pages)


def _parse_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _parse_pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        texts.append(text)
    return "\n\n".join(texts)


def _parse_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    texts = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        texts.append(f"Sheet: {sheet}")
        for row in ws.iter_rows(values_only=True):
            row_text = " | ".join(str(cell) for cell in row if cell is not None)
            if row_text:
                texts.append(row_text)
    return "\n".join(texts)
